"""Unit and integration tests for document processing handler."""

import os
import shutil
import tempfile
import unittest
from docx import Document
from docx.shared import Pt, RGBColor

from backend.app.doc_handler import (
    process_document,
    process_docx_document,
    merge_identical_adjacent_runs,
    run_libreoffice_conversion,
)


class TestDocHandler(unittest.TestCase):

    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_typecast_")

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_run_fragmentation_merging(self):
        """Tests that split digraphs across adjacent identical runs are merged cleanly."""
        doc = Document()
        p = doc.add_paragraph()

        # Simulate Word splitting "Shahar" into 3 runs: "S", "h", "ahar"
        r1 = p.add_run("S")
        r2 = p.add_run("h")
        r3 = p.add_run("ahar")

        # Set matching styling
        for r in (r1, r2, r3):
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(12)

        self.assertEqual(len(p.runs), 3)

        # Merge runs
        merge_identical_adjacent_runs(p)

        self.assertEqual(len(p.runs), 1)
        self.assertEqual(p.runs[0].text, "Shahar")
        self.assertTrue(p.runs[0].bold)

    def test_docx_in_place_transliteration(self):
        """Creates a comprehensive docx with headers, footers, tables, and runs, then verifies transliteration."""
        input_docx = os.path.join(self.test_dir, "sample_input.docx")
        output_docx = os.path.join(self.test_dir, "sample_output.docx")

        doc = Document()

        # 1. Fragmented paragraph: "O", "'", "zbekiston"
        p1 = doc.add_paragraph()
        r1 = p1.add_run("O")
        r2 = p1.add_run("'")
        r3 = p1.add_run("zbekiston Respublikasi")
        for r in (r1, r2, r3):
            r.font.name = "Calibri"

        # 2. Table
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).paragraphs[0].text = "Shahar"
        table.cell(0, 1).paragraphs[0].text = "Toshkent"
        table.cell(1, 0).paragraphs[0].text = "Viloyat"
        table.cell(1, 1).paragraphs[0].text = "Samarqand"

        # 3. Header
        header = doc.sections[0].header
        header.paragraphs[0].text = "Rasmiy hujjat"

        doc.save(input_docx)

        # Transliterate docx
        process_docx_document(input_docx, output_docx, direction="latin-to-cyrillic")

        # Verify output
        res_doc = Document(output_docx)
        p_text = res_doc.paragraphs[0].text
        self.assertEqual(p_text, "Ўзбекистон Республикаси")

        # Verify table cells
        self.assertEqual(res_doc.tables[0].cell(0, 0).paragraphs[0].text, "Шаҳар")
        self.assertEqual(res_doc.tables[0].cell(0, 1).paragraphs[0].text, "Тошкент")
        self.assertEqual(res_doc.tables[0].cell(1, 0).paragraphs[0].text, "Вилоят")
        self.assertEqual(res_doc.tables[0].cell(1, 1).paragraphs[0].text, "Самарқанд")

        # Verify header
        self.assertEqual(res_doc.sections[0].header.paragraphs[0].text, "Расмий ҳужжат")

    def test_txt_transliteration(self):
        """Tests plain text document conversion."""
        input_txt = os.path.join(self.test_dir, "test.txt")
        with open(input_txt, "w", encoding="utf-8") as f:
            f.write("O'zbekiston kelajagi buyuk davlat.\nEshikni yoping.")

        out_dir = os.path.join(self.test_dir, "out_txt")
        result_path = process_document(input_txt, out_dir, direction="latin-to-cyrillic", target_format="txt")

        with open(result_path, "r", encoding="utf-8") as f:
            content = f.read()

        self.assertIn("Ўзбекистон келажаги буюк давлат.", content)
        self.assertIn("Эшикни ёпинг.", content)

    def test_libreoffice_pdf_export(self):
        """Tests LibreOffice conversion to PDF if LibreOffice is present."""
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            self.skipTest("LibreOffice not installed on system.")

        input_txt = os.path.join(self.test_dir, "doc_test.txt")
        with open(input_txt, "w", encoding="utf-8") as f:
            f.write("Toshkent shahri - O'zbekiston poytaxti.")

        out_dir = os.path.join(self.test_dir, "out_pdf")
        pdf_path = process_document(input_txt, out_dir, direction="latin-to-cyrillic", target_format="pdf")

        self.assertTrue(os.path.exists(pdf_path))
        self.assertTrue(pdf_path.endswith(".pdf"))
        self.assertGreater(os.path.getsize(pdf_path), 0)

    def test_doc_auto_modernization_to_docx(self):
        """Tests that .doc files default to modern .docx output under Option 3."""
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            self.skipTest("LibreOffice not installed on system.")

        input_txt = os.path.join(self.test_dir, "modernize.txt")
        with open(input_txt, "w", encoding="utf-8") as f:
            f.write("O'zbekiston - buyuk yurt.")

        doc_path = run_libreoffice_conversion(input_txt, "doc", self.test_dir)
        self.assertTrue(doc_path.endswith(".doc"))

        # Under Option 3, target_format="same" modernizes .doc to .docx
        result = process_document(
            doc_path,
            os.path.join(self.test_dir, "out_modern"),
            direction="latin-to-cyrillic",
            target_format="same",
        )
        self.assertTrue(result.endswith(".docx"))


if __name__ == "__main__":
    unittest.main()
