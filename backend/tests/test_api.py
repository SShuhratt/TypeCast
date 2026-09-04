"""Integration tests for FastAPI REST endpoints."""

import io
import unittest
from fastapi.testclient import TestClient
from docx import Document

from backend.app.main import app

client = TestClient(app)


class TestAPIEndpoints(unittest.TestCase):

    def test_health_check(self):
        response = client.get("/api/health")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["status"], "healthy")
        self.assertTrue(data["libreoffice_available"])

    def test_text_transliterate_endpoint(self):
        payload = {
            "text": "O'zbekiston - go'zal yurt!",
            "direction": "latin-to-cyrillic"
        }
        response = client.post("/api/transliterate/text", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["transliterated"], "Ўзбекистон - гўзал юрт!")
        self.assertEqual(data["direction"], "latin-to-cyrillic")

    def test_text_transliterate_reverse_endpoint(self):
        payload = {
            "text": "Ўзбекистон - гўзал юрт!",
            "direction": "cyrillic-to-latin"
        }
        response = client.post("/api/transliterate/text", json=payload)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["transliterated"], "O'zbekiston - go'zal yurt!")

    def test_file_upload_txt(self):
        file_content = "Salom dunyo! O'zbekiston kelajagi buyuk davlat.".encode("utf-8")
        files = {
            "file": ("test.txt", io.BytesIO(file_content), "text/plain")
        }
        data = {
            "direction": "latin-to-cyrillic",
            "target_format": "txt"
        }
        response = client.post("/api/transliterate/file", files=files, data=data)
        self.assertEqual(response.status_code, 200)
        content = response.content.decode("utf-8")
        self.assertIn("Салом дунё! Ўзбекистон келажаги буюк давлат.", content)

    def test_file_upload_docx(self):
        # Generate in-memory docx
        doc = Document()
        doc.add_paragraph("Shahar va viloyatlar")
        doc_stream = io.BytesIO()
        doc.save(doc_stream)
        doc_stream.seek(0)

        files = {
            "file": ("doc_sample.docx", doc_stream, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        }
        data = {
            "direction": "latin-to-cyrillic",
            "target_format": "docx"
        }
        response = client.post("/api/transliterate/file", files=files, data=data)
        self.assertEqual(response.status_code, 200)

        # Parse returned docx
        out_doc = Document(io.BytesIO(response.content))
        self.assertEqual(out_doc.paragraphs[0].text, "Шаҳар ва вилоятлар")


if __name__ == "__main__":
    unittest.main()
